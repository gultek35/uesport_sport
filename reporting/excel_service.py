"""
Excel Rapor Servisi - Sporcu performans raporunu Excel olarak olusturur.
"""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime


def generate_athlete_report_excel(report_data):
    """
    Sporcu performans raporunu Excel olarak olusturur.

    Args:
        report_data: Rapor verileri (dict)

    Returns:
        BytesIO: Excel dosyasi
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Performans Raporu"

    # 📌 Stiller
    header_font = Font(name="Arial", size=12, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1a1a2e", end_color="1a1a2e", fill_type="solid")
    header_alignment = Alignment(horizontal="center", vertical="center")

    cell_font = Font(name="Arial", size=10)
    cell_alignment = Alignment(horizontal="center", vertical="center")

    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin")
    )

    row = 1

    # 📌 Baslik
    ws.merge_cells(f"A{row}:F{row}")
    title_cell = ws.cell(row=row, column=1, value="UESPORT - Performans Raporu")
    title_cell.font = Font(name="Arial", size=16, bold=True, color="1a1a2e")
    title_cell.alignment = Alignment(horizontal="center")
    row += 2

    # 📌 Sporcu Bilgileri
    athlete = report_data.get("sporcu", {})
    ws.cell(row=row, column=1, value="Sporcu:").font = Font(bold=True)
    ws.cell(row=row, column=2, value=athlete.get("ad_soyad", ""))
    ws.cell(row=row, column=4, value="Yas:").font = Font(bold=True)
    ws.cell(row=row, column=5, value=athlete.get("yas", ""))
    row += 1

    ws.cell(row=row, column=1, value="Cinsiyet:").font = Font(bold=True)
    ws.cell(row=row, column=2, value=athlete.get("cinsiyet", ""))
    ws.cell(row=row, column=4, value="Seviye:").font = Font(bold=True)
    ws.cell(row=row, column=5, value=athlete.get("level", ""))
    row += 1

    ws.cell(row=row, column=1, value="Rapor Tarihi:").font = Font(bold=True)
    ws.cell(row=row, column=2, value=datetime.now().strftime("%d.%m.%Y %H:%M"))
    row += 2

    # 📌 Test Sonuclari Tablosu Basligi
    ws.cell(row=row, column=1, value="📋 Test Sonuclari").font = Font(name="Arial", size=12, bold=True)
    row += 1

    # 📌 Tablo Basliklari
    headers = ["Test", "Deger", "Yuzdelik Dilim", "Z-Skor", "Durum", "AI Yorum"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=row, column=col, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_alignment
        cell.border = border
    row += 1

    # 📌 Veriler
    results = report_data.get("sonuclar", [])
    for r in results:
        norm = r.get("norm_bilgisi", {})

        ws.cell(row=row, column=1, value=r.get("test_adi", "")).border = border
        ws.cell(row=row, column=2, value=f"{r.get('deger', '')} {r.get('birim', '')}").border = border
        ws.cell(row=row, column=3, value=f"{norm.get('yuzdelik_dilim', '')}%" if norm else "-").border = border
        ws.cell(row=row, column=4, value=norm.get("z_skor", "-") if norm else "-").border = border
        ws.cell(row=row, column=5, value=norm.get("etiket", "-") if norm else "-").border = border
        ws.cell(row=row, column=6, value=r.get("ai_yorum", "")).border = border

        for col in range(1, 7):
            ws.cell(row=row, column=col).alignment = cell_alignment
            ws.cell(row=row, column=col).font = cell_font

        # Durum renklendirme
        label = norm.get("etiket", "") if norm else ""
        if label in ["Mükemmel", "Çok yi", "yi"]:
            ws.cell(row=row, column=5).fill = PatternFill(start_color="d4edda", end_color="d4edda", fill_type="solid")
        elif label in ["Zayıf", "Çok Zayıf"]:
            ws.cell(row=row, column=5).fill = PatternFill(start_color="f8d7da", end_color="f8d7da", fill_type="solid")
        else:
            ws.cell(row=row, column=5).fill = PatternFill(start_color="fff3cd", end_color="fff3cd", fill_type="solid")

        row += 1

    row += 1

    # 📌 Ozet
    ozet = report_data.get("ozet", {})
    ws.cell(row=row, column=1, value="📝 Performans Ozeti").font = Font(name="Arial", size=12, bold=True)
    row += 1
    ws.cell(row=row, column=1, value=ozet.get("genel_degerlendirme", "")).font = Font(size=10)
    ws.merge_cells(f"A{row}:F{row}")
    row += 2

    # 📌 Güçlü Yönler ve Gelisim Alanlari
    strong = ozet.get("guclu_yonler", [])
    weak = ozet.get("gelisim_alanlari", [])

    if strong:
        ws.cell(row=row, column=1, value="💪 Güçlü Yönler:").font = Font(bold=True, color="28a745")
        ws.cell(row=row, column=2, value=", ".join(strong))
        row += 1

    if weak:
        ws.cell(row=row, column=1, value="📈 Gelisim Alanlari:").font = Font(bold=True, color="fd7e14")
        ws.cell(row=row, column=2, value=", ".join(weak))
        row += 1

    row += 1

    # 📌 AI Degerlendirmesi
    ai_ozet = report_data.get("ai_ozet", "")
    if ai_ozet:
        ws.cell(row=row, column=1, value="🤖 AI Degerlendirmesi").font = Font(name="Arial", size=12, bold=True, color="6f42c1")
        row += 1
        ws.cell(row=row, column=1, value=ai_ozet).font = Font(size=10)
        ws.merge_cells(f"A{row}:F{row}")

    # 📌 Sutun Genislikleri
    column_widths = [25, 15, 15, 12, 15, 50]
    for i, width in enumerate(column_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width

    # 📌 Kaydet
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
